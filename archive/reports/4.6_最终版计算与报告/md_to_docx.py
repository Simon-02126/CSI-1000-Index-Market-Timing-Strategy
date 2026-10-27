# -*- coding: utf-8 -*-
"""把 中证1000择时策略说明_最终版.md 转成排版良好的 .docx"""
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import os as _os_path, sys as _sys_path
_d = _os_path.path.dirname(_os_path.path.abspath(__file__))
while _d != _os_path.path.dirname(_d) and not _os_path.path.isfile(
        _os_path.path.join(_d, "_paths.py")):
    _d = _os_path.path.dirname(_d)
_sys_path.path.insert(0, _d)
import _paths as _P  # noqa: E402  统一路径入口

HERE = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = _P.REPORTS                  # md / png / docx 三者同目录
MD = os.path.join(SRC_DIR, "中证1000择时策略说明_最终版.md")
OUT = os.path.join(SRC_DIR, "中证1000择时策略说明_最终版.docx")

HAN_HEAD = "微软雅黑"
HAN_BODY = "宋体"
LATIN_BODY = "Calibri"
MONO = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)

TOKEN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+?\*)")


# ---------- 底层工具 ----------
def style_run(run, size=10.5, bold=False, italic=False, mono=False,
              han=HAN_BODY, latin=LATIN_BODY, color=None):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = MONO if mono else latin
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), MONO if mono else latin)
    rFonts.set(qn("w:hAnsi"), MONO if mono else latin)
    rFonts.set(qn("w:eastAsia"), han)
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size=10.5, bold=False, han=HAN_BODY, latin=LATIN_BODY,
               color=None, italic=False):
    """解析 **粗体** / `代码` / *斜体* 并写入段落（斜体内部可再嵌行内语法）"""
    text = text.replace("\\|", "|")
    for seg in TOKEN.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            r = par.add_run(seg[2:-2])
            style_run(r, size=size, bold=True, italic=italic, han=han, latin=latin, color=color)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            body = seg[1:-1]
            has_cjk = any(ord(c) > 0x2E80 for c in body)
            r = par.add_run(body)
            style_run(r, size=size - 0.5, mono=not has_cjk, italic=italic,
                      han=HAN_BODY, latin=(HAN_BODY if has_cjk else LATIN_BODY),
                      color=RGBColor(0xA3, 0x15, 0x15))
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            add_inline(par, seg[1:-1], size=size, bold=bold, han=han, latin=latin,
                       color=color, italic=True)
        else:
            r = par.add_run(seg)
            style_run(r, size=size, bold=bold, italic=italic, han=han, latin=latin, color=color)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hex_fill)
    tcPr.append(el)


def cell_margins(table, top=30, bottom=30, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{tag}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def disp_len(s):
    return sum(2 if ord(c) > 0x2E80 else 1 for c in s)


def fix_table_width(tbl, rows, ncol, total_cm=17.6):
    """按各列内容长度按比例分配列宽，并用固定布局锁定"""
    w = []
    for c in range(ncol):
        m = max(disp_len(r[c]) for r in rows)
        w.append(min(max(m, 10), 42))
    tot = float(sum(w))
    widths = [Cm(total_cm * x / tot) for x in w]
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(Cm(total_cm).twips)))
    tblW.set(qn("w:type"), "dxa")
    for ci, cw in enumerate(widths):
        tbl.columns[ci].width = cw
        for r in tbl.rows:
            r.cells[ci].width = cw


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    style_run(run, size=9, color=GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._element.addnext(fld)


# ---------- 文档骨架 ----------
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(1.7)
sec.right_margin = Cm(1.7)
CONTENT_W = Cm(21.0 - 1.7 * 2)

normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.font.name = LATIN_BODY
normal.element.rPr.rFonts.set(qn("w:eastAsia"), HAN_BODY)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.28

doc.core_properties.title = "中证1000 择时策略说明（最终版）"
doc.core_properties.subject = "A股宽基指数多因子择时"
doc.core_properties.comments = "量价 + 资金面 + 情绪面 三块合成择时，2019-06-04 ~ 2026-09-30"

add_page_number_footer(sec)

# ---------- 逐行解析 ----------
lines = open(MD, encoding="utf-8").read().split("\n")
i = 0
first_h1 = True
in_code = False
code_buf = []


def flush_code():
    global code_buf
    if not code_buf:
        return
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = CONTENT_W
    cell = tbl.cell(0, 0)
    cell.width = CONTENT_W
    shade(cell, "F5F5F5")
    cell_margins(tbl, top=60, bottom=60, left=120, right=120)
    p0 = cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(0)
    p0.paragraph_format.line_spacing = 1.0
    for k, cl in enumerate(code_buf):
        par = p0 if k == 0 else cell.add_paragraph()
        if k:
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.0
        r = par.add_run(cl if cl else " ")
        style_run(r, size=8.5, mono=True)
    code_buf = []
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


while i < len(lines):
    raw = lines[i]
    line = raw.rstrip()

    # 代码围栏
    if line.startswith("```"):
        if in_code:
            flush_code()
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    # 空行
    if not line.strip():
        i += 1
        continue

    # 分隔线
    if re.fullmatch(r"-{3,}", line.strip()):
        i += 1
        continue

    # 图片
    m = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
    if m:
        path = os.path.join(SRC_DIR, m.group(2))
        if os.path.exists(path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.add_run().add_picture(path, width=Cm(17.0))
        i += 1
        continue

    # 标题
    m = re.match(r"^(#{1,4}) (.+)$", line)
    if m:
        lvl = len(m.group(1))
        text = m.group(2).strip()
        if lvl == 1 and first_h1:
            first_h1 = False
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(14)
            add_inline(p, text, size=22, bold=True, han=HAN_HEAD, latin=HAN_HEAD, color=ACCENT)
        else:
            size = {1: 16, 2: 13, 3: 11.5, 4: 10.5}[lvl]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt({1: 16, 2: 12, 3: 9, 4: 7}[lvl])
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_inline(p, text, size=size, bold=True, han=HAN_HEAD, latin=HAN_HEAD, color=ACCENT)
            if lvl == 1:
                pPr = p._p.get_or_add_pPr()
                bdr = OxmlElement("w:pBdr")
                bot = OxmlElement("w:bottom")
                bot.set(qn("w:val"), "single")
                bot.set(qn("w:sz"), "8")
                bot.set(qn("w:space"), "4")
                bot.set(qn("w:color"), "1F4E79")
                bdr.append(bot)
                pPr.append(bdr)
        i += 1
        continue

    # 表格
    if line.strip().startswith("|") and i + 1 < len(lines) and re.match(
            r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i].strip())
            i += 1
        rows = []
        for b in block:
            if re.match(r"^\|[\s:|-]+\|$", b):
                continue
            inner = b.strip("|")
            inner = inner.replace("\\|", "\x00")
            cells = [c.replace("\x00", "|").strip() for c in inner.split("|")]
            rows.append(cells)
        if not rows:
            continue
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        fs = 9.0 if ncol <= 4 else (8.5 if ncol <= 6 else 7.5)
        tbl = doc.add_table(rows=len(rows), cols=ncol)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True
        cell_margins(tbl, top=25, bottom=25, left=60, right=60)
        for ri, r in enumerate(rows):
            for ci, txt in enumerate(r):
                cell = tbl.cell(ri, ci)
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.line_spacing = 1.05
                if ri == 0:
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    add_inline(par, txt, size=fs, bold=True, han=HAN_HEAD, latin=HAN_HEAD)
                else:
                    add_inline(par, txt, size=fs)
        for c in tbl.rows[0].cells:
            shade(c, "DCE6F1")
        repeat_header(tbl.rows[0])
        fix_table_width(tbl, rows, ncol)
        for r in tbl.rows:
            trPr = r._tr.get_or_add_trPr()
            ns = OxmlElement("w:cantSplit")
            ns.set(qn("w:val"), "true")
            trPr.append(ns)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue

    # 引用块
    if line.strip().startswith(">"):
        buf = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            buf.append(lines[i].strip().lstrip(">").strip())
            i += 1
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = CONTENT_W
        cell = tbl.cell(0, 0)
        cell.width = CONTENT_W
        shade(cell, "F2F6FA")
        cell_margins(tbl, top=70, bottom=70, left=140, right=140)
        for k, q in enumerate([x for x in buf if x]):
            par = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.2
            add_inline(par, q, size=9.5, han=HAN_BODY, color=RGBColor(0x33, 0x33, 0x33))
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        continue

    # 列表
    m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.+)$", line)
    if m:
        indent = len(m.group(1))
        ordered = bool(re.match(r"\d+\.", m.group(2)))
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65 + 0.55 * (indent // 2))
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        bullet = m.group(2) if ordered else ("·" if indent < 2 else "–")
        r = p.add_run(bullet + "  ")
        style_run(r, size=10.5, han=HAN_BODY, color=ACCENT)
        add_inline(p, m.group(3), size=10.5)
        i += 1
        continue

    # 普通段落
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.0)
    add_inline(p, line.strip(), size=10.5)
    i += 1

if in_code:
    flush_code()

doc.save(OUT)
print("saved:", OUT)

# ---------- 复核 ----------
chk = Document(OUT)
print("段落数:", len(chk.paragraphs), " 表格数:", len(chk.tables))
print("内联图片数:", len(chk.inline_shapes))
heads = [p.text for p in chk.paragraphs if p.text.strip() and p.runs and p.runs[0].font.size
         and p.runs[0].font.size.pt >= 13]
print("主标题/大标题：")
for h in heads[:12]:
    print("   ", h[:60])
words = sum(len(p.text) for p in chk.paragraphs)
tw = sum(len(c.text) for t in chk.tables for r in t.rows for c in r.cells)
print(f"正文字数 {words}，表格字数 {tw}，合计 {words + tw}")
